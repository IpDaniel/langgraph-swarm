import requests, os, time, io
import pandas as pd
import xlrd
from bs4 import BeautifulSoup
from docx import Document
from dotenv import load_dotenv
from src.agent.schemas.file import StateFile

load_dotenv()

class Normalizer:
   def __init__(self):
      self.dispatch_table = {
            "csv": self._handle_csv,
            "xls": self._handle_xls,
            "xlsx": self._handle_xlsx,
            "txt": self._handle_txt,
            "pdf": self._handle_pdf,
            "html": self._handle_html,
            "docx": self._handle_docx,
         }
   
   def normalize(self, file: StateFile) -> str:
      suffix = file.title.split('.')[-1].lower()
      return self.dispatch_table[suffix](file)

   def _handle_csv(self, f: StateFile) -> str:
      csv_bytes = f.as_bytes()
      return csv_bytes.decode('utf-8')

   def _handle_xls(self, f: StateFile) -> str:
      file_bytes = f.as_bytes()
      workbook = xlrd.open_workbook(file_contents=file_bytes)
      
      content_parts = []
      for sheet in workbook.sheets():
         content_parts.append(f"\n=== Sheet: {sheet.name} ===\n")
         data = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
         if data:
               df = pd.DataFrame(data[1:], columns=data[0]) if len(data) > 1 else pd.DataFrame()
               # Convert float columns that are actually integers
               for col in df.columns:
                  if df[col].dtype == 'float64':
                     # Check if all non-null values are whole numbers
                     if df[col].notna().any() and (df[col].dropna() % 1 == 0).all():
                           df[col] = df[col].astype('Int64')  # Nullable integer type
               content_parts.append(df.to_csv(index=False))
      
      return "\n".join(content_parts)

   def _handle_xlsx(self, f: StateFile) -> str:
      file_bytes = f.as_bytes()
      bytes_io = io.BytesIO(file_bytes)
      try:
         excel_file = pd.ExcelFile(bytes_io, engine="openpyxl")
         loader = lambda name: pd.read_excel(excel_file, sheet_name=name)
         sheet_names = excel_file.sheet_names
      except Exception:
         # Fallback: some .xlsx samples may actually be legacy/invalid zips; try xlrd
         workbook = xlrd.open_workbook(file_contents=file_bytes)
         sheet_names = [s.name for s in workbook.sheets()]
         def loader(name):
               sheet = workbook.sheet_by_name(name)
               data = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
               if not data:
                  return pd.DataFrame()
               return pd.DataFrame(data[1:], columns=data[0]) if len(data) > 1 else pd.DataFrame()
      
      content_parts = []
      for sheet_name in sheet_names:
         df: pd.DataFrame = loader(sheet_name)
         
         # Convert float columns that are actually integers
         for col in df.columns:
               if df[col].dtype == 'float64':
                  # Check if all non-null values are whole numbers
                  if df[col].notna().any() and (df[col].dropna() % 1 == 0).all():
                     df[col] = df[col].astype('Int64')  # Nullable integer type
         
         content_parts.append(f"\n=== Sheet: {sheet_name} ===\n")
         content_parts.append(df.to_csv(index=False))
      
      return "\n".join(content_parts)

   def _handle_txt(self, f: StateFile) -> str:
      txt_bytes = f.as_bytes()
      return txt_bytes.decode('utf-8')

   def _handle_pdf(self, f: StateFile) -> str:
      BASE = "https://api.mathpix.com/v3"
      HEADERS = {"app_id": os.getenv("MATHPIX_APP_ID"), "app_key": os.getenv("MATHPIX_APP_KEY")}

      pdf_bytes = f.as_bytes()
      fileobj = io.BytesIO(pdf_bytes)
      files = {"file": (f.title or "upload.pdf", fileobj, "application/pdf")}
      data = {"options_json": '{"conversion_formats":{"md":true}}'}

      r = requests.post(f"{BASE}/pdf", headers=HEADERS, files=files, data=data, timeout=120)
      r.raise_for_status()
      pdf_id = r.json()["pdf_id"]

      status_url = f"{BASE}/pdf/{pdf_id}"
      while True:
         s = requests.get(status_url, headers=HEADERS, timeout=30).json()
         if s.get("status") == "completed":
               break
         if s.get("status") in {"error", "failed"}:
               raise SystemExit(s)
         time.sleep(1)

      md = requests.get(f"{status_url}.md", headers=HEADERS, timeout=300).text
      return md

   def _handle_html(self, f: StateFile) -> str:
      html_bytes = f.as_bytes()
      html_content = html_bytes.decode('utf-8')
      soup = BeautifulSoup(html_content, 'html.parser')
      
      for script in soup(["script", "style"]):
         script.decompose()
      
      text = soup.get_text()
      lines = (line.strip() for line in text.splitlines())
      chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
      text = '\n'.join(chunk for chunk in chunks if chunk)
      return text

   def _handle_docx(self, f: StateFile) -> str:
      docx_bytes = f.as_bytes()
      bytes_io = io.BytesIO(docx_bytes)
      doc = Document(bytes_io)
      
      paragraphs = []
      for para in doc.paragraphs: # paragraphs
         if para.text.strip():
               paragraphs.append(para.text)
      
      for table in doc.tables: # tables
         for row in table.rows:
               row_data = []
               for cell in row.cells:
                  if cell.text.strip():
                     row_data.append(cell.text.strip())
               if row_data:
                  paragraphs.append(" | ".join(row_data))
      
      return "\n".join(paragraphs)